
from pathlib import Path
import sqlite3
import pandas as pd

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter


# 1. Project paths
BASE_DIR = Path(__file__).resolve().parent.parent
DB_PATH = BASE_DIR / "data" / "fintech_loan_analysis.db"
SQL_DIR = BASE_DIR / "sql"
OUTPUT_PATH = BASE_DIR / "reports" / "fintech_management_report.xlsx"


# 2. SQL files and worksheet names
REPORTS = [
    ("summary.sql", "Overall Summary"),
    ("contract_type_analysis.sql", "Contract Analysis"),
    ("credit_income_analysis.sql", "Credit-Income Analysis"),
    ("income_type_analysis.sql", "Income Type Analysis"),
]


# 3. Key business observations
findings = [
    [
        "Overall recorded difficulty rate",
        "24,825 out of 307,511 applications (8.07%)",
        "This is the overall recorded repayment-difficulty share.",
        "This is a historical dataset outcome, not a prediction of future defaults."
    ],
    [
        "Contract type comparison",
        "Cash loans: 8.35%; Revolving loans: 5.48%",
        "Cash loans have a higher observed difficulty rate in this dataset.",
        "This comparison is descriptive and does not establish causation."
    ],
    [
        "Credit-to-income comparison",
        "The 2–4x band has the highest observed rate: 8.87%.",
        "The observed rates vary across credit-to-income bands.",
        "The pattern is not monotonic; these bands are exploratory, not lending thresholds."
    ],
    [
        "Income type comparison",
        "Working: 9.59%; State servant: 5.75%; Pensioner: 5.39%.",
        "Observed rates differ across income categories.",
        "Very small groups, such as Maternity leave (5 applicants), have unstable rates."
    ],
]


def load_query(connection, filename):
    """Read a SQL file and return its query results as a DataFrame."""
    query_path = SQL_DIR / filename

    if not query_path.exists():
        raise FileNotFoundError(f"SQL file not found: {query_path}")

    query = query_path.read_text(encoding="utf-8")
    return pd.read_sql_query(query, connection)


def format_worksheet(ws):
    """Apply consistent formatting to a worksheet."""
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )
    header_font = Font(
        bold=True,
        color="FFFFFF"
    )
    thin_border = Border(
        bottom=Side(style="thin", color="D9E2F3")
    )

    # Format column headers (row 4)
    for cell in ws[4]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    # Format data rows
    for row in ws.iter_rows(min_row=5):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True
            )

            # Format percentage columns
            header = ws.cell(row=4, column=cell.column).value
            if header and "pct" in str(header).lower():
                cell.number_format = '0.00'

    # Adjust column widths
    for column_cells in ws.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = 0

        for cell in column_cells:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))

        ws.column_dimensions[column_letter].width = min(
            max(max_length + 4, 16),
            42
        )

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[4].height = 32


def write_dataframe(writer, dataframe, sheet_name, title):
    """Write a DataFrame and add a title and formatting."""
    dataframe.to_excel(
        writer,
        sheet_name=sheet_name,
        index=False,
        startrow=3
    )

    ws = writer.sheets[sheet_name]

    ws["A1"] = title
    ws["A1"].font = Font(
        size=16,
        bold=True,
        color="1F4E78"
    )

    ws["A2"] = (
        "Source: Home Credit application-level dataset | "
        "Historical descriptive analysis"
    )
    ws["A2"].font = Font(
        italic=True,
        color="666666"
    )

    format_worksheet(ws)

    # Add a chart when a difficulty-rate column is available
    headers = list(dataframe.columns)

    rate_columns = [
        index + 1
        for index, name in enumerate(headers)
        if "difficulty_rate_pct" in str(name).lower()
    ]

    if rate_columns and len(dataframe) > 0:
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = "Recorded Repayment-Difficulty Rate"
        chart.y_axis.title = "Category"
        chart.x_axis.title = "Rate (%)"
        chart.height = 7
        chart.width = 14

        data = Reference(
            ws,
            min_col=rate_columns[0],
            min_row=4,
            max_row=4 + len(dataframe)
        )
        categories = Reference(
            ws,
            min_col=1,
            min_row=5,
            max_row=4 + len(dataframe)
        )

        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        ws.add_chart(chart, "G4")


def main():
    if not DB_PATH.exists():
        raise FileNotFoundError(
            f"Database not found: {DB_PATH}"
        )

    print("Connecting to SQLite database...")

    with sqlite3.connect(DB_PATH) as connection:
        with pd.ExcelWriter(
            OUTPUT_PATH,
            engine="openpyxl"
        ) as writer:

            # Create the four analysis worksheets
            for sql_file, sheet_name in REPORTS:
                print(f"Preparing {sheet_name}...")

                df = load_query(connection, sql_file)

                write_dataframe(
                    writer,
                    df,
                    sheet_name,
                    sheet_name
                )

            # Create the Key Findings worksheet
            findings_df = pd.DataFrame(
                findings,
                columns=[
                    "Business Finding",
                    "Evidence",
                    "Interpretation",
                    "Caveat"
                ]
            )

            write_dataframe(
                writer,
                findings_df,
                "Key Findings",
                "Key Business Findings"
            )

    print("\nExcel report created successfully!")
    print(f"Saved at: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
